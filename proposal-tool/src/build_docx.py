"""
Step 3: render a proposal dict into a .docx draft.
"""

from pathlib import Path

from docx import Document


def build_docx(company_name: str, proposal: dict, out_path: Path) -> None:
    doc = Document()

    doc.add_heading(f"Proposal Draft — {company_name}", level=1)

    doc.add_heading("Situation", level=2)
    doc.add_paragraph(proposal.get("situation_summary") or "(not generated)")

    doc.add_heading("Recommended Services", level=2)
    services = proposal.get("recommended_services") or []
    if not services:
        doc.add_paragraph("(no services recommended)")
    for svc in services:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(
            f"{svc.get('service_name', 'Unnamed service')} — {svc.get('indicative_price', 'TBD')}"
        )
        run.bold = True
        doc.add_paragraph(svc.get("rationale", ""))

    doc.add_heading("Next Step", level=2)
    doc.add_paragraph(proposal.get("next_step") or "(not generated)")

    gaps = proposal.get("gaps_to_confirm") or []
    if gaps:
        doc.add_heading("Gaps to confirm before sending", level=2)
        for gap in gaps:
            doc.add_paragraph(gap, style="List Bullet")

    if "raw_output" in proposal:
        doc.add_heading("Raw model output (JSON parse failed)", level=2)
        doc.add_paragraph(proposal["raw_output"])

    doc.save(str(out_path))


def safe_filename(company_name: str) -> str:
    """Collapse a company name to a filesystem-safe slug; never empty."""
    slug = "".join(c if c.isalnum() else "_" for c in company_name).strip("_")
    return slug or "unnamed_company"


def unique_output_path(directory: Path, base_name: str, date_str: str, suffix: str = ".docx") -> Path:
    """Avoid silently overwriting an existing draft for the same company/day."""
    candidate = directory / f"{base_name}_{date_str}{suffix}"
    if not candidate.exists():
        return candidate

    n = 2
    while True:
        candidate = directory / f"{base_name}_{date_str}_{n}{suffix}"
        if not candidate.exists():
            return candidate
        n += 1
